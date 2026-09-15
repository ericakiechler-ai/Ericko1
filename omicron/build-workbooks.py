#!/usr/bin/env python3
"""Turn the generated test-plan CSVs into one Excel workbook per relay.

The CSVs are the machine-readable form; these are the form a technician
carries onto the job — the steps sheet is laid out to be filled in, with
Measured and Pass/Fail as the only cells anyone should type into.

The values are NOT retyped here either. This script reads the CSVs that
build-testplan.js produced, so the chain from the tool's own calculation
through to the workbook is unbroken.

Usage:  python3 build-workbooks.py
"""
import csv, os, re, textwrap
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'xlsx')

RELAYS = [
    ('845-VEC', 'Multilin 845 transformer protection', ['845-VEC'], '845-TestUniverse-build-guide.md'),
    ('850-FDR', 'Multilin 850 feeder protection', ['850-FDR'], '850-TestUniverse-build-guide.md'),
    ('869-MTR', 'Multilin 869 motor protection', ['869-MTR'], '869-TestUniverse-build-guide.md'),
    ('889-GEN', 'Multilin 889 generator protection', ['889-GEN'], '889-TestUniverse-build-guide.md'),
    ('7SJ85-SIP', 'Siemens 7SJ85 SIPROTEC 5 overcurrent', ['7SJ85-SIP'], '7SJ85-TestUniverse-build-guide.md'),
    ('7SD82-DIF', 'Siemens 7SD82 SIPROTEC 5 line differential',
     ['7SD82-DIF', '7SD82-DIF-loopback'], '7SD82-TestUniverse-build-guide.md'),
]

FONT = 'Arial'
HEAD_FILL = PatternFill('solid', fgColor='2A2E31')
FILLIN_FILL = PatternFill('solid', fgColor='FFF2A8')   # cells to type into
BAND_FILL = PatternFill('solid', fgColor='F2F4F5')
WARN_FILL = PatternFill('solid', fgColor='FBE2E4')
TITLE_FILL = PatternFill('solid', fgColor='E2E7E9')
THIN = Side(style='thin', color='C9CFD2')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# header -> (display name, width, alignment)
COLS = {
    'Element':        ('Function', 10, 'left'),
    'ElementName':    ('Description', 26, 'left'),
    'TU_Module':      ('TU module', 17, 'left'),
    'DriveAs':        ('Drive as', 16, 'left'),
    'Step':           ('#', 5, 'center'),
    'StepName':       ('Test step', 34, 'left'),
    'WhatIsApplied':  ('What is applied', 52, 'left'),
    'SeqRun':         ('Run', 5, 'center'),
    'SeqState':       ('State', 7, 'center'),
    'StateDuration_s': ('Dur s', 7, 'center'),
    'V_L1': ('V L1', 15, 'center'), 'V_L2': ('V L2', 15, 'center'),
    'V_L3': ('V L3', 15, 'center'), 'V_4': ('V 4', 15, 'center'),
    'I_L1': ('I L1', 15, 'center'), 'I_L2': ('I L2', 15, 'center'),
    'I_L3': ('I L3', 15, 'center'),
    'I_B_L1': ('I B-L1', 15, 'center'), 'I_B_L2': ('I B-L2', 15, 'center'),
    'I_B_L3': ('I B-L3', 15, 'center'),
    'Freq_Hz': ('Hz', 8, 'center'),
    'CurrentNeededAt': ('Current at', 11, 'center'),
    'Reachable': ('Reachable', 26, 'left'),
    'Expected': ('Expected', 20, 'left'),
    'ExpectedDetail': ('Expected — detail', 40, 'left'),
    'Measured': ('MEASURED', 15, 'center'),
    'PassFail': ('PASS / FAIL', 12, 'center'),
    # channel-sheet extras
    'Group': ('Group', 16, 'left'),
    'SeqLabel': ('State label', 20, 'left'),
    'Channel': ('Channel', 11, 'left'),
    'Magnitude': ('Magnitude', 12, 'center'),
    'Unit': ('Unit', 6, 'center'),
    'Angle_deg': ('Angle deg', 10, 'center'),
    'Frequency_Hz': ('Hz', 8, 'center'),
    'ChannelNote': ('Note', 46, 'left'),
    'Verdict': ('Verdict', 8, 'center'),
}
FILLIN = {'Measured', 'PassFail'}


def read_csv(path):
    """Split the leading '#' comment block from the table."""
    notes, rows = [], []
    with open(path, newline='', encoding='utf-8') as f:
        raw = f.read().split('\n')
    body = []
    for line in raw:
        if line.startswith('#'):
            notes.append(line.lstrip('# ').rstrip())
        else:
            body.append(line)
    rdr = csv.reader(body)
    for r in rdr:
        if r and any(c.strip() for c in r):
            rows.append(r)
    return notes, rows[0], rows[1:]


def drop_empty(header, rows):
    """A column that is blank on every row is noise on a printed sheet —
       except the two the technician is meant to fill in."""
    keep = [i for i, h in enumerate(header)
            if h in FILLIN or any((r[i] if i < len(r) else '').strip() for r in rows)]
    return [header[i] for i in keep], [[(r[i] if i < len(r) else '') for i in keep] for r in rows]


def style_header(ws, header, row=1):
    for c, h in enumerate(header, 1):
        disp, width, _ = COLS.get(h, (h, 14, 'left'))
        cell = ws.cell(row=row, column=c, value=disp)
        cell.font = Font(name=FONT, bold=True, size=9, color='FFFFFF')
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BOX
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[row].height = 28


def write_table(ws, header, rows, start=1, band_key=0):
    style_header(ws, header, start)
    prev, band = None, False
    for i, r in enumerate(rows):
        key = r[band_key] if band_key < len(r) else ''
        if key != prev:
            band = not band
            prev = key
        for c, h in enumerate(header, 1):
            v = r[c - 1]
            cell = ws.cell(row=start + 1 + i, column=c)
            cell.value = float(v) if _numeric(h, v) else (v or None)
            _, _, al = COLS.get(h, (h, 14, 'left'))
            cell.font = Font(name=FONT, size=9)
            cell.alignment = Alignment(horizontal=al, vertical='top', wrap_text=(al == 'left'))
            cell.border = BOX
            if h in FILLIN:
                cell.fill = FILLIN_FILL
            elif h == 'Reachable' and v.startswith('NO'):
                cell.fill = WARN_FILL
                cell.font = Font(name=FONT, size=9, color='9B2C2C', bold=True)
            elif band:
                cell.fill = BAND_FILL
    ws.freeze_panes = ws.cell(row=start + 1, column=3)
    ws.auto_filter.ref = f'A{start}:{get_column_letter(len(header))}{start + len(rows)}'


NUMCOLS = {'Step', 'Magnitude', 'Angle_deg', 'Frequency_Hz', 'Freq_Hz',
           'StateDuration_s', 'SeqRun', 'Group'}


def _numeric(h, v):
    if h not in NUMCOLS or h == 'Group' or not v:
        return False
    try:
        float(v); return True
    except ValueError:
        return False


def add_validation(ws, header, nrows, start=1):
    if 'PassFail' not in header:
        return
    col = get_column_letter(header.index('PassFail') + 1)
    dv = DataValidation(type='list', formula1='"PASS,FAIL,N/A,NOT TESTED"', allow_blank=True)
    dv.prompt = 'Select a result'
    dv.promptTitle = 'Result'
    ws.add_data_validation(dv)
    dv.add(f'{col}{start + 1}:{col}{start + nrows}')


def label_notes(notes):
    """The CSV header block is prose. Give each line the label a reader of a
       test record would expect, rather than a column of 'Note'."""
    out, pending = [], None
    for n in notes:
        if not n:
            continue
        if n.startswith('configuration:'):
            out.append(('Configuration', n.split(':', 1)[1].strip()))
        elif n.lower().startswith('generated'):
            out.append(('Generated', n))
        elif ' — ' in n and not out:
            out.append(('Plan', n))
        elif n.startswith('All quantities'):
            out.append(('Conventions', n))
        elif n.startswith('Values shown'):
            out.append(('Sample data', n))
        elif n.startswith('Staging'):
            pending = ('Staging', n)
        elif pending:
            pending = (pending[0], pending[1] + ' ' + n)
        else:
            out.append(('Note', n))
    if pending:
        out.append(pending)
    return out


def setup_sheet(wb, tag, title, notes, sheets, pf_col):
    """Front sheet: what produced the plan, how to use it, and a live count
       of what has been signed off — the counts are formulas so they move as
       the sheet is filled in."""
    ws = wb.create_sheet('Start here', 0)
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 96
    r = 1
    ws.merge_cells('A1:B1')
    c = ws['A1']; c.value = f'{tag} — {title}'
    c.font = Font(name=FONT, bold=True, size=14)
    c.fill = TITLE_FILL
    c.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 30
    r = 3

    def block(head, lines, fillin=False):
        nonlocal r
        ws.cell(row=r, column=1, value=head).font = Font(name=FONT, bold=True, size=10)
        ws.cell(row=r, column=1).fill = TITLE_FILL
        ws.cell(row=r, column=2).fill = TITLE_FILL
        r += 1
        for k, v in lines:
            ws.cell(row=r, column=1, value=k).font = Font(name=FONT, size=9, bold=True)
            ws.cell(row=r, column=1).alignment = Alignment(vertical='top')
            cc = ws.cell(row=r, column=2, value=v or None)
            cc.font = Font(name=FONT, size=9)
            cc.alignment = Alignment(wrap_text=True, vertical='top')
            if fillin:
                cc.fill = FILLIN_FILL
                cc.border = BOX
            ws.row_dimensions[r].height = max(14, 12 * (1 + len(v) // 105))
            r += 1
        r += 1

    block('How to use this workbook', [
        ('Fill in', 'Only the two shaded columns on each test sheet: MEASURED and PASS / FAIL. '
                    'Everything else is calculated and should not be edited by hand.'),
        ('Example row', 'MEASURED "1.432 s", PASS / FAIL "PASS". Use the same unit the Expected '
                        'column uses, so the two can be compared without conversion.'),
        ('Untested steps', 'Record them as NOT TESTED with a reason. A blank result reads as work '
                           'not yet done; it must never be left blank on a completed record.'),
        ('If a setting changes', 'Do not edit numbers here. Change it in the relay tool, re-run the '
                                 'generator, and rebuild this workbook — that is what keeps the '
                                 'injected values and the calculations from drifting apart.'),
    ])

    block('Where these values come from', label_notes(notes))

    block('Record identification', [
        ('Station', ''), ('Feeder / plant item', ''), ('Device serial', ''),
        ('Firmware', ''), ('Setting file revision', ''), ('Test set serial', ''),
        ('Tested by', ''), ('Date', ''), ('Witnessed by', ''),
    ], fillin=True)

    block('Sheets in this workbook', [(nm, desc) for nm, desc in sheets])

    # live progress counters — formulas, so they follow the sheet as it fills
    first = sheets[0][0]
    ws.cell(row=r, column=1, value='Progress').font = Font(name=FONT, bold=True, size=10)
    ws.cell(row=r, column=1).fill = TITLE_FILL
    ws.cell(row=r, column=2).fill = TITLE_FILL
    r += 1
    q = f"'{first}'" if ' ' in first else first
    rng = f'{q}!${pf_col}$2:${pf_col}$2000'
    for lab, formula in [
        ('Total steps', f'=COUNTA({q}!$A$2:$A$2000)'),
        ('Results entered', f'=COUNTA({rng})'),
        ('Passed', f'=COUNTIF({rng},"PASS")'),
        ('Failed', f'=COUNTIF({rng},"FAIL")'),
        ('Not tested', f'=COUNTIF({rng},"NOT TESTED")'),
        ('Outstanding', f'=COUNTA({q}!$A$2:$A$2000)-COUNTA({rng})'),
    ]:
        ws.cell(row=r, column=1, value=lab).font = Font(name=FONT, size=9, bold=True)
        cc = ws.cell(row=r, column=2, value=formula)
        cc.font = Font(name=FONT, size=9)
        cc.alignment = Alignment(horizontal='left')
        r += 1
    ws.cell(row=r, column=2, value='Counts refer to the first test sheet only.').font = \
        Font(name=FONT, size=8, italic=True, color='6B7477')
    return ws


def guide_sheet(wb, path):
    if not os.path.exists(path):
        return
    ws = wb.create_sheet('Build guide')
    ws.column_dimensions['A'].width = 118
    r = 1
    for line in open(path, encoding='utf-8').read().split('\n'):
        s = line.rstrip()
        cell = ws.cell(row=r, column=1)
        if s.startswith('#'):
            lvl = len(s) - len(s.lstrip('#'))
            cell.value = s.lstrip('# ')
            cell.font = Font(name=FONT, bold=True, size=13 - lvl)
            cell.fill = TITLE_FILL if lvl <= 2 else PatternFill()
        elif s.startswith('|'):
            cell.value = s
            cell.font = Font(name='Courier New', size=8)
        else:
            cell.value = s
            cell.font = Font(name=FONT, size=9)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = max(13, 12 * (1 + len(s) // 118))
        r += 1
    return ws


def build(tag, title, plans, guide):
    wb = Workbook()
    wb.remove(wb.active)
    sheets, notes0, pf_col = [], None, 'A'
    for plan in plans:
        suffix = plan[len(tag):].lstrip('-')
        step_name = 'Test steps' + (f' ({suffix})' if suffix else '')
        chan_name = 'Channels' + (f' ({suffix})' if suffix else '')
        notes, hdr, rows = read_csv(os.path.join(HERE, plan + '-steps.csv'))
        notes0 = notes0 or notes
        hdr, rows = drop_empty(hdr, rows)
        ws = wb.create_sheet(step_name)
        write_table(ws, hdr, rows)
        add_validation(ws, hdr, len(rows))
        if 'PassFail' in hdr and pf_col == 'A':
            pf_col = get_column_letter(hdr.index('PassFail') + 1)
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = '1:1'
        sheets.append((step_name, f'{len(rows)} test steps — the record. Fill in MEASURED and PASS / FAIL.'))

        cnotes, chdr, crows = read_csv(os.path.join(HERE, plan + '-channels.csv'))
        chdr, crows = drop_empty(chdr, crows)
        cws = wb.create_sheet(chan_name)
        write_table(cws, chdr, crows)
        cws.page_setup.orientation = 'landscape'
        sheets.append((chan_name, f'{len(crows)} rows — one per channel per step, for pasting '
                                  f'into a Test Universe module table.'))
    if guide:
        sheets.append(('Build guide', 'How to assemble this plan in Test Universe, and the thing '
                                      'about this relay most likely to be got wrong.'))
    setup_sheet(wb, tag, title, notes0, sheets, pf_col)
    guide_sheet(wb, os.path.join(HERE, guide))
    out = os.path.join(OUT, f'{tag}-test-plan.xlsx')
    wb.save(out)
    return out, sum(1 for _ in wb.sheetnames)


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    for tag, title, plans, guide in RELAYS:
        path, n = build(tag, title, plans, guide)
        print(f'{os.path.basename(path):32s} {n} sheets')
